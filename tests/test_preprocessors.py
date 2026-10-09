"""Unit tests for preprocessors.recover_unplaced_images."""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx
import pytest
from docx.shared import Inches

import preprocessors

Image = pytest.importorskip("PIL.Image")


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), "white").save(buf, format="PNG")
    return buf.getvalue()


PNG = _png()


def _docx_with_picture(path: Path) -> None:
    d = docx.Document()
    d.add_paragraph("Body text.")
    d.add_picture(io.BytesIO(PNG), width=Inches(1))
    d.save(str(path))


def test_image_missing_from_markdown_is_saved_and_reported(tmp_path):
    src = tmp_path / "source.docx"
    _docx_with_picture(src)
    (tmp_path / "article.md").write_text("Body text.\n", encoding="utf-8")

    recovered = preprocessors.recover_unplaced_images(src, tmp_path)

    assert recovered == ["assets/media/image1.png"]
    assert (tmp_path / "assets" / "media" / "image1.png").read_bytes() == PNG


def test_image_already_in_markdown_is_not_reported(tmp_path):
    src = tmp_path / "source.docx"
    _docx_with_picture(src)
    (tmp_path / "article.md").write_text(
        "Body text.\n\n![](assets/media/image1.png)\n", encoding="utf-8"
    )
    assert preprocessors.recover_unplaced_images(src, tmp_path) == []


def test_unreadable_input_reports_nothing(tmp_path):
    (tmp_path / "article.md").write_text("Body.\n", encoding="utf-8")
    bogus = tmp_path / "not-a-docx.docx"
    bogus.write_bytes(b"not a zip")
    assert preprocessors.recover_unplaced_images(bogus, tmp_path) == []
